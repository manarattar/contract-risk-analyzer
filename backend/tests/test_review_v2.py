"""Offline safety regression tests; only synthetic temporary records are used."""
import hashlib
import io
import json
import time
import zipfile

import pytest
from fastapi.testclient import TestClient

from app.review.settings import ReviewSettings, settings
from app.review.store import initialize, transaction
from app.review import jobs, pipeline
from app.review.parser import ParseError, parse
from app.review.ingestion import object_path

TOKENS = {"alice": "a"*48, "bob": "b"*48, "viewer": "v"*48}


@pytest.fixture
def client(tmp_path, monkeypatch):
    principals = [{"id": name, "workspace_id": "workspace-a" if name != "bob" else "workspace-b",
                   "role": "viewer" if name == "viewer" else "reviewer",
                   "token_sha256": hashlib.sha256(token.encode()).hexdigest()} for name, token in TOKENS.items()]
    monkeypatch.setenv("REVIEW_MODE", "demo")
    monkeypatch.setenv("REVIEW_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("REVIEW_PRINCIPALS_JSON", json.dumps(principals))
    monkeypatch.setenv("REVIEW_UPLOAD_ENABLED", "false")
    settings.cache_clear()
    initialize()
    from app.main import app
    def no_provider(*args, **kwargs):
        raise AssertionError("Provider calls are forbidden in offline tests")
    monkeypatch.setattr(pipeline, "generate", no_provider)
    monkeypatch.setattr(jobs, "parse_in_process", lambda d: parse(object_path(d["id"], d["kind"]), d["kind"]))
    with TestClient(app) as result:
        yield result
    settings.cache_clear()


def auth(user="alice"):
    return {"Authorization": "Bearer " + TOKENS[user]}


def sample(client, key="sample-key-00000001"):
    response = client.post("/api/v2/sample", headers=auth() | {"Idempotency-Key": key})
    assert response.status_code == 202, response.text
    return response.json()["id"]


def processed(client):
    identifier = sample(client)
    assert jobs.run_once()
    result = client.get(f"/api/v2/documents/{identifier}", headers=auth())
    assert result.status_code == 200
    assert result.json()["status"] == "ready", result.text
    return identifier, result.json()


@pytest.mark.parametrize("mode", ["disabled", "manual", "demo"])
def test_no_implicit_live_or_mock(mode):
    cfg = ReviewSettings(mode=mode, _env_file=None)
    assert cfg.mode == mode
    assert cfg.upload_enabled is False


def test_live_and_upload_gates():
    with pytest.raises(ValueError):
        ReviewSettings(mode="live", api_key="fake", model="fake", _env_file=None)
    with pytest.raises(ValueError):
        ReviewSettings(mode="manual", upload_enabled=True, _env_file=None)
    with pytest.raises(ValueError):
        ReviewSettings(allowed_origins="*", _env_file=None)


def test_anonymous_and_legacy_routes(client):
    assert client.get("/api/v2/documents").status_code == 401
    for path in ["/api/status/x", "/api/analysis/x", "/api/report/x", "/api/compare/result/x"]:
        assert client.get(path, headers=auth()).status_code == 404
    for path in ["/api/upload", "/api/qa", "/api/compare"]:
        assert client.post(path, headers=auth()).status_code == 404


@pytest.mark.parametrize("method,suffix,body", [
    ("get", "", None), ("get", "/source", None), ("get", "/export", None),
    ("post", "/questions", {"question": "What are payment terms?"}),
    ("post", "/retry", None), ("post", "/cancel", None),
    ("delete", "", None),
    ("patch", "/findings/sample-liability", {"status": "accepted", "version": 0}),
    ("post", "/complete", {"revision": 0, "acknowledge_incomplete": True}),
])
def test_foreign_document_denied(client, method, suffix, body):
    identifier = sample(client)
    kwargs = {"headers": auth("bob")}
    if body is not None:
        kwargs["json"] = body
    response = getattr(client, method)(f"/api/v2/documents/{identifier}{suffix}", **kwargs)
    assert response.status_code == 404, response.text
    assert "Synthetic services" not in response.text


def test_foreign_comparison_and_library(client):
    identifier = sample(client)
    response = client.post("/api/v2/comparisons", headers=auth("bob"), json={"baseline_id": identifier, "revised_id": identifier})
    assert response.status_code == 404
    assert client.get("/api/v2/documents", headers=auth("bob")).json()["items"] == []


def test_viewer_cannot_write(client):
    assert client.post("/api/v2/sample", headers=auth("viewer") | {"Idempotency-Key": "viewer-key-0000001"}).status_code == 403


def test_idempotency_no_duplicate_file_or_job(client):
    first = sample(client)
    assert sample(client) == first
    assert len(list((settings().data_dir / "uploads").iterdir())) == 1
    with transaction() as conn:
        assert conn.execute("SELECT count(*) FROM jobs").fetchone()[0] == 1


def test_source_provenance_and_no_canned_answer(client):
    identifier, doc = processed(client)
    assert doc["result"]["provenance"]["mode"] == "demo"
    citation = doc["result"]["findings"][0]["citations"][0]
    source = client.get(f"/api/v2/documents/{identifier}/source?block_id={citation['block_id']}", headers=auth()).json()
    block = source["blocks"][0]
    assert block["text"][citation["start"]:citation["end"]] == citation["quote"]
    answer = client.post(f"/api/v2/documents/{identifier}/questions", headers=auth(), json={"question": "astronaut zebras galaxies?"}).json()
    assert answer["status"] == "abstained" and answer["sources"] == []
    assert "15 days" not in answer["answer"]
    legal = client.post(f"/api/v2/documents/{identifier}/questions", headers=auth(), json={"question": "Is it safe to sign?"}).json()
    assert legal["status"] == "abstained"


def test_decision_conflict_and_escaped_export(client):
    identifier, _ = processed(client)
    endpoint = f"/api/v2/documents/{identifier}/findings/sample-liability"
    payload = {"status": "edited", "note": '<script>alert("x")</script>', "revision_text": '<img src="https://evil.invalid/a">', "version": 0}
    assert client.patch(endpoint, headers=auth(), json=payload).status_code == 200
    assert client.patch(endpoint, headers=auth(), json=payload).status_code == 409
    response = client.get(f"/api/v2/documents/{identifier}/export", headers=auth())
    assert response.status_code == 200
    assert "<script>" not in response.text and "&lt;script&gt;" in response.text
    assert '<img src="https://' not in response.text
    assert response.headers["cache-control"] == "no-store"
    assert "Incomplete human review" in response.text
    with transaction() as conn:
        assert conn.execute("SELECT count(*) FROM decision_events").fetchone()[0] == 1


def test_completion_requires_decisions_and_revision(client):
    identifier, _ = processed(client)
    endpoint = f"/api/v2/documents/{identifier}"
    payload = {"revision": 0, "acknowledge_incomplete": True}
    assert client.post(endpoint+"/complete", headers=auth(), json=payload).status_code == 409
    assert client.patch(endpoint+"/findings/sample-liability", headers=auth(), json={"status": "accepted", "version": 0}).status_code == 200
    assert client.post(endpoint+"/complete", headers=auth(), json=payload).status_code == 409
    payload["revision"] = 1
    assert client.post(endpoint+"/complete", headers=auth(), json=payload).status_code == 200


def test_cancel_invalidates_late_job_and_retry(client):
    identifier = sample(client)
    job = jobs.claim()
    assert client.post(f"/api/v2/documents/{identifier}/cancel", headers=auth()).status_code == 200
    assert not jobs.checkpoint(job, "complete", result={"coverage": "readable"}, final=True)
    assert client.post(f"/api/v2/documents/{identifier}/retry", headers=auth()).status_code == 202
    assert jobs.run_once()
    assert client.get(f"/api/v2/documents/{identifier}", headers=auth()).json()["status"] == "ready"


def test_expired_lease_recovery(client):
    identifier = sample(client)
    original = jobs.claim()
    with transaction() as conn:
        conn.execute("UPDATE jobs SET lease_until=? WHERE document_id=?", (time.time()-1, identifier))
    replacement = jobs.claim()
    assert replacement["lease_token"] != original["lease_token"]
    assert not jobs.checkpoint(original, "late")
    assert jobs.checkpoint(replacement, "recovered")


def test_parser_failure_never_low_risk(client, monkeypatch):
    identifier = sample(client)
    def fail(_doc):
        raise ValueError("no_readable_text")
    monkeypatch.setattr(jobs, "parse_in_process", fail)
    jobs.run_once()
    doc = client.get(f"/api/v2/documents/{identifier}", headers=auth()).json()
    assert doc["status"] == "failed" and doc["result"] is None
    assert "risk_score" not in json.dumps(doc)


def test_deletion_revokes_purges_and_prevents_resurrection(client):
    identifier, _ = processed(client)
    assert client.delete(f"/api/v2/documents/{identifier}", headers=auth()).status_code == 202
    assert client.get(f"/api/v2/documents/{identifier}", headers=auth()).status_code == 404
    assert client.get(f"/api/v2/documents/{identifier}/export", headers=auth()).status_code == 404
    assert client.get(f"/api/v2/deletions/{identifier}", headers=auth("bob")).status_code == 404
    jobs.sweep()
    assert not object_path(identifier, "txt").exists()
    receipt = client.get(f"/api/v2/deletions/{identifier}", headers=auth()).json()
    assert receipt["state"] == "complete"
    assert client.delete(f"/api/v2/documents/{identifier}", headers=auth()).status_code == 202
    assert not jobs.run_once()
    with transaction() as conn:
        row = conn.execute("SELECT source,result,context,filename FROM documents WHERE id=?", (identifier,)).fetchone()
        assert row["source"] is None and row["result"] is None and row["context"] == "{}"


def test_read_only_source_search_and_compare(client):
    identifier, _ = processed(client)
    response = client.post("/api/v2/comparisons", headers=auth(), json={"baseline_id": identifier, "revised_id": identifier})
    assert response.status_code == 200 and response.json()["changes"] == []
    assert "winner" not in response.json()


def test_ingestion_gate_size_and_signature(client, monkeypatch):
    endpoint = "/api/v2/documents"
    fields = {"context": '{"party":"Example"}', "privacy_acknowledged": "true"}
    headers = auth() | {"Idempotency-Key": "upload-key-0000001"}
    assert client.post(endpoint, headers=headers, data=fields, files={"file": ("a.txt", b"Synthetic text")}).status_code == 403
    monkeypatch.setenv("REVIEW_MODE", "manual")
    monkeypatch.setenv("REVIEW_UPLOAD_ENABLED", "true")
    monkeypatch.setenv("REVIEW_PARSER_ISOLATION_APPROVED", "true")
    monkeypatch.setenv("REVIEW_MAX_BYTES", "100")
    settings.cache_clear()
    response = client.post(endpoint, headers=headers, data=fields, files={"file": ("a.pdf", b"not a pdf")})
    assert response.status_code == 415
    assert client.post(endpoint, headers=headers, data=fields, files={"file": ("a.txt", b"x"*101)}).status_code == 413
    response = client.post(endpoint, headers=headers, data=fields, files={"file": ("../../safe.txt", b"Synthetic clause.")})
    assert response.status_code == 202, response.text
    assert response.json()["filename"] == "safe.txt"
    jobs.run_once()
    assert all(p.parent == settings().data_dir / "uploads" for p in (settings().data_dir / "uploads").iterdir())


def test_preserves_short_preamble_and_long_sentence(tmp_path):
    text = "Preamble.\n\n1. IP\nOurs.\n\n" + "x"*11000
    path = tmp_path / "test.txt"
    path.write_text(text, encoding="utf-8")
    result = parse(path, "txt")
    normalized = "".join(b["text"] for b in result["blocks"])
    assert "Preamble." in normalized and "Ours." in normalized
    assert normalized.count("x") == 11000
    assert max(len(b["text"]) for b in result["blocks"]) <= 5000


def test_human_annotation_has_owner_source_and_event(client):
    identifier, _ = processed(client)
    payload = {"block_id": "b1", "title": "Check the source", "note": "Synthetic reviewer observation."}
    endpoint = f"/api/v2/documents/{identifier}/annotations"
    assert client.post(endpoint, headers=auth("bob"), json=payload).status_code == 404
    payload["block_id"] = "foreign"
    assert client.post(endpoint, headers=auth(), json=payload).status_code == 404
    payload["block_id"] = "b1"
    assert client.post(endpoint, headers=auth(), json=payload).status_code == 201
    doc = client.get(f"/api/v2/documents/{identifier}", headers=auth()).json()
    assert doc["result"]["findings"][-1]["evidence_status"] == "Human annotation by alice"


def test_evaluation_requires_real_adjudication():
    from app.review.evaluation import measures, release_gate
    metrics = measures(["a", "a", "wrong"], ["a", "b"], 1, 2, 1, 2)
    assert metrics["finding_precision"] == .5 and metrics["finding_recall"] == .5
    assert measures([], ["expected"], 0, 0, 0, 0)["finding_recall"] == 0
    assert release_gate({})["eligible"] is False


def test_failed_generation_export_remains_incomplete(client):
    identifier = sample(client)
    with transaction() as conn:
        conn.execute("UPDATE documents SET mode='live' WHERE id=?", (identifier,))
    # The fixture forbids provider use; failure is a terminal incomplete state.
    jobs.run_once()
    doc = client.get(f"/api/v2/documents/{identifier}", headers=auth()).json()
    assert doc["status"] == "failed"
    assert doc["result"]["coverage"] == "partial"
    exported = client.get(f"/api/v2/documents/{identifier}/export", headers=auth())
    assert "processing did not complete" in exported.text
    assert "No standard clauses are missing" not in exported.text


def test_changed_wording_not_hidden_by_scores(client):
    first, _ = processed(client)
    second = sample(client, "second-synthetic-001")
    jobs.run_once()
    with transaction() as conn:
        row = conn.execute("SELECT source FROM documents WHERE id=?", (second,)).fetchone()
        source = json.loads(row[0])
        block = next(b for b in source["blocks"] if "preceding thirty days" in b["text"])
        block["text"] = block["text"].replace("preceding thirty days", "preceding twelve months")
        conn.execute("UPDATE documents SET source=? WHERE id=?", (json.dumps(source), second))
    result = client.post("/api/v2/comparisons", headers=auth(), json={"baseline_id": first, "revised_id": second}).json()
    assert result["changes"] and result["changes"][0]["kind"] == "Changed"
    assert "timing" in result["changes"][0]["assessment"]
    assert "winner" not in result


def test_stream_limit_without_content_length(client, monkeypatch):
    monkeypatch.setenv("REVIEW_MODE", "manual")
    monkeypatch.setenv("REVIEW_UPLOAD_ENABLED", "true")
    monkeypatch.setenv("REVIEW_PARSER_ISOLATION_APPROVED", "true")
    monkeypatch.setenv("REVIEW_MAX_BYTES", "100")
    settings.cache_clear()
    boundary = "syntheticboundary"
    prefix = f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="a.txt"\r\nContent-Type: text/plain\r\n\r\n'.encode()
    def body():
        yield prefix
        for _ in range(20):
            yield b"x" * 8192
        yield f"\r\n--{boundary}--\r\n".encode()
    response = client.post("/api/v2/documents", headers=auth() | {"Idempotency-Key": "stream-key-00000001", "Content-Type": f"multipart/form-data; boundary={boundary}"}, content=body())
    assert response.status_code == 413, response.text


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -1, 1.1, True, "0.99"])
def test_evaluation_gate_rejects_invalid_metrics(value):
    from app.review.evaluation import release_gate
    report = {"reviewer_adjudicated": True, "qualified_reviewers": ["reviewer-a", "reviewer-b"],
              "held_out_families": 40, "dataset_hash": "synthetic", "prompt_hash": "synthetic",
              "model_version": "synthetic", "parser_version": "synthetic", "critical_failures": 0,
              "metrics": {"citation_accuracy": 1, "finding_precision": 1, "finding_recall": 1, "unsupported_claim_rate": 0}}
    assert release_gate(report)["eligible"]
    report["held_out_families"] = value
    assert not release_gate(report)["eligible"]
    report["held_out_families"] = 40
    report["metrics"]["citation_accuracy"] = value
    assert not release_gate(report)["eligible"]
    report["metrics"]["citation_accuracy"] = 1
    report["metrics"]["unsupported_claim_rate"] = value
    assert not release_gate(report)["eligible"]
    report["metrics"]["unsupported_claim_rate"] = 0
    report["qualified_reviewers"] = ["reviewer-a", "reviewer-a"]
    assert not release_gate(report)["eligible"]


def test_docx_tables_and_archive_safety(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Synthetic body.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Synthetic table obligation."
    path = tmp_path / "table.docx"
    doc.save(path)
    result = parse(path, "docx")
    assert any("table obligation" in b["text"] for b in result["blocks"])
    bad = tmp_path / "bad.docx"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("../escaped", "synthetic")
    with pytest.raises(ParseError):
        parse(bad, "docx")


@pytest.mark.parametrize("change", ["missing", "reordered", "wrong_quote", "wrong_block", "assurance"])
def test_rejects_invalid_generated_evidence(change):
    blocks = [{"id": "b1", "text": "Payment within fifteen days.", "location": "Paragraph 1"}, {"id": "b2", "text": "Synthetic other clause.", "location": "Paragraph 2"}]
    raw = {"block_ids": ["b1", "b2"], "findings": [{"id": "f1", "title": "Check payment timing", "explanation": "Verify the intended deadline.", "uncertainty": "Business preference unknown", "action": "Ask the reviewer.", "citations": [{"block_id": "b1", "quote": "Payment within fifteen days."}]}]}
    if change == "missing": raw["block_ids"] = ["b1"]
    if change == "reordered": raw["block_ids"].reverse()
    if change == "wrong_quote": raw["findings"][0]["citations"][0]["quote"] = "Payment in one year."
    if change == "wrong_block": raw["findings"][0]["citations"][0]["block_id"] = "foreign"
    if change == "assurance": raw["findings"][0]["action"] = "Safe to sign."
    with pytest.raises(ValueError):
        pipeline.validate_generated(raw, blocks)
